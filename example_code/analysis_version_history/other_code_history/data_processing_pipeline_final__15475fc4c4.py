#!/usr/bin/env python3
"""Final scoped Part 1 pipeline for the haemograph blood-rheology cohort.

This copy preserves the teammate pipeline and adds explicit exclusion of DataZ and
Excel lock files when selecting workbooks. It does not calculate LAOS harmonics,
wavelets, model fits, or feature-selection results. Source workbooks are read only
and are never saved or overwritten.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import shutil
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


ARROW_RE = re.compile(
    r"(?P<a>[+-]?(?:\d+(?:\.\d*)?|\.\d+))\s*(?:-->|-\s*>|→|\bto\b)\s*"
    r"(?P<b>[+-]?(?:\d+(?:\.\d*)?|\.\d+))",
    re.I,
)
GAMMA_RE = re.compile(r"(?:gamma|amp(?:litude)?)\s*=\s*([0-9]*\.?[0-9]+)", re.I)
OMEGA_RE = re.compile(r"(?:omega|ω)\s*=\s*([0-9]*\.?[0-9]+)", re.I)


def norm(value) -> str:
    return re.sub(r"\s+", " ", "" if value is None else str(value).strip()).lower()


def number(value) -> float:
    if value is None or isinstance(value, bool):
        return math.nan
    try:
        out = float(value)
    except (TypeError, ValueError):
        return math.nan
    return out if math.isfinite(out) else math.nan


def sheet_type(name: str) -> str:
    s = norm(name)
    if "udlaos" in s:
        return "UDLAOS"
    if "laos" in s:
        return "LAOS"
    if "cessation" in s:
        return "CESSATION"
    if "reversal" in s or "reverse" in s:
        return "REVERSAL"
    if "step" in s or re.search(r"\d\s+to\s+[+-]?\d", s):
        return "STEP_UP_DOWN"
    if "helton" in s or "srfs" in s or "sawtooth" in s:
        return "OTHER_TRANSIENT"
    if "saos" in s or "steady" in s or s == "ss" or s.startswith("ss ") or s.startswith(" ss"):
        return "SS_SAOS"
    if "amplitude sweep" in s or "frequency sweep" in s:
        return "SS_SAOS"
    if "donor" in s or "lab data" in s:
        return "ASPIRIN_SUMMARY"
    return "OTHER"


def parse_condition(label: str, protocol: str) -> dict:
    out = {
        "step_start_rate_label_1_s": math.nan,
        "step_final_rate_label_1_s": math.nan,
        "step_magnitude_label_1_s": math.nan,
        "step_direction": "",
        "omega_label_rad_s": math.nan,
        "strain_amplitude_label_fraction": math.nan,
        "strain_amplitude_label_pct": math.nan,
    }
    if protocol == "LAOS":
        g = GAMMA_RE.search(label)
        w = OMEGA_RE.search(label)
        if g:
            gamma = float(g.group(1))
            out["strain_amplitude_label_fraction"] = gamma
            out["strain_amplitude_label_pct"] = 100.0 * gamma
        if w:
            out["omega_label_rad_s"] = float(w.group(1))
    else:
        m = ARROW_RE.search(label)
        if m:
            start = float(m.group("a"))
            final = float(m.group("b"))
            out["step_start_rate_label_1_s"] = start
            out["step_final_rate_label_1_s"] = final
            out["step_magnitude_label_1_s"] = abs(final - start)
            out["step_direction"] = "UP" if final > start else "DOWN" if final < start else "UNKNOWN"
    return out


def nearest_label(ws, header_row: int, col: int) -> str:
    for row in range(header_row - 1, 0, -1):
        for c in range(col, max(0, col - 3), -1):
            value = ws.cell(row, c).value
            if value not in (None, ""):
                return str(value).strip()
    return ws.title


def trace_hash(frame: pd.DataFrame) -> str:
    vals = frame[["time_s", "stress_pa", "strain_pct", "shear_rate_1_s"]].to_numpy(float)
    vals = np.nan_to_num(vals, nan=9.87654321e99, posinf=8.7654321e99, neginf=-8.7654321e99)
    return hashlib.sha256(np.round(vals, 8).tobytes()).hexdigest()


def inspect_sheet(ws) -> dict:
    nonempty = numeric = formulas = 0
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if value is None:
                continue
            nonempty += 1
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                numeric += 1
            if isinstance(value, str) and value.startswith("="):
                formulas += 1
    stype = sheet_type(ws.title)
    return {
        "sheet": ws.title,
        "sheet_type": stype,
        "rows": ws.max_row,
        "columns": ws.max_column,
        "nonempty_cells": nonempty,
        "numeric_cells": numeric,
        "formula_cells": formulas,
        "part1_core_scope": "YES" if stype in {"LAOS", "STEP_UP_DOWN"} else "NO",
        "scope_note": "Core LAOS/step-up/down raw time series" if stype in {"LAOS", "STEP_UP_DOWN"} else "Catalogued but not included in the Part 1 core long table",
    }


def extract_core_blocks(path: Path) -> tuple[list[dict], list[pd.DataFrame], list[dict], list[dict]]:
    wb = load_workbook(path, read_only=False, data_only=True)
    measurements: list[dict] = []
    signals: list[pd.DataFrame] = []
    issues: list[dict] = []
    exclusions: list[dict] = []
    for ws in wb.worksheets:
        stype = sheet_type(ws.title)
        if stype not in {"LAOS", "STEP_UP_DOWN"}:
            exclusions.append(
                {
                    "source_file": path.name,
                    "sheet": ws.title,
                    "sheet_type": stype,
                    "target": "worksheet",
                    "reason": "Outside Part 1 core scope; retained in sheet inventory",
                }
            )
            continue
        headers = []
        for row in range(1, min(ws.max_row, 12) + 1):
            for col in range(1, ws.max_column + 1):
                if norm(ws.cell(row, col).value).startswith("step time"):
                    headers.append((row, col))
        if not headers:
            exclusions.append(
                {
                    "source_file": path.name,
                    "sheet": ws.title,
                    "sheet_type": stype,
                    "target": "worksheet",
                    "reason": "No Step time header detected",
                }
            )
            continue
        for header_row, col in headers:
            next_cols = [c for r, c in headers if r == header_row and c > col]
            stop_col = min(next_cols) if next_cols else min(ws.max_column + 1, col + 8)
            fields = {"time_s": col}
            for c in range(col + 1, stop_col):
                h = norm(ws.cell(header_row, c).value)
                if "stress" in h and "stress_pa" not in fields:
                    fields["stress_pa"] = c
                elif "strain" in h and "strain_pct" not in fields:
                    fields["strain_pct"] = c
                elif ("shear rate" in h or h == "shear") and "shear_rate_1_s" not in fields:
                    fields["shear_rate_1_s"] = c
            label = nearest_label(ws, header_row, col)
            if "stress_pa" not in fields:
                exclusions.append(
                    {
                        "source_file": path.name,
                        "sheet": ws.title,
                        "sheet_type": stype,
                        "target": f"block column {col}",
                        "reason": f"Stress channel missing ({label})",
                    }
                )
                continue
            if stype == "LAOS" and "strain_pct" not in fields:
                exclusions.append(
                    {
                        "source_file": path.name,
                        "sheet": ws.title,
                        "sheet_type": stype,
                        "target": f"block column {col}",
                        "reason": f"Strain channel missing ({label})",
                    }
                )
                continue
            start_row = header_row + 1
            while start_row <= ws.max_row and not math.isfinite(number(ws.cell(start_row, col).value)):
                start_row += 1
            rows = []
            empty_run = 0
            for row in range(start_row, ws.max_row + 1):
                t = number(ws.cell(row, fields["time_s"]).value)
                stress = number(ws.cell(row, fields["stress_pa"]).value)
                if not (math.isfinite(t) and math.isfinite(stress)):
                    empty_run += 1
                    if empty_run > 40 and len(rows) >= 10:
                        break
                    continue
                empty_run = 0
                strain = number(ws.cell(row, fields["strain_pct"]).value) if "strain_pct" in fields else math.nan
                rate = number(ws.cell(row, fields["shear_rate_1_s"]).value) if "shear_rate_1_s" in fields else math.nan
                rows.append((row, t, stress, strain, rate))
            if len(rows) < 12:
                exclusions.append(
                    {
                        "source_file": path.name,
                        "sheet": ws.title,
                        "sheet_type": stype,
                        "target": f"block column {col}",
                        "reason": f"Fewer than 12 aligned numeric rows ({label})",
                    }
                )
                continue
            frame = pd.DataFrame(rows, columns=["source_row", "time_s", "stress_pa", "strain_pct", "shear_rate_1_s"])
            parsed = parse_condition(label, stype)
            dt = np.diff(frame.time_s.to_numpy())
            positive_dt = dt[dt > 0]
            median_dt = float(np.median(positive_dt)) if len(positive_dt) else math.nan
            dt_cv = float(np.std(positive_dt) / np.mean(positive_dt)) if len(positive_dt) > 1 and np.mean(positive_dt) else math.nan
            measured_start = measured_final = math.nan
            if frame.shear_rate_1_s.notna().sum() >= 12:
                n = len(frame)
                measured_start = float(frame.shear_rate_1_s.iloc[: max(3, int(0.03 * n))].median())
                measured_final = float(frame.shear_rate_1_s.iloc[-max(5, int(0.2 * n)) :].median())
                if not parsed["step_direction"] and stype == "STEP_UP_DOWN":
                    change = measured_final - measured_start
                    if abs(change) > 0.05 * max(abs(measured_start), abs(measured_final), 0.1):
                        parsed["step_direction"] = "UP" if change > 0 else "DOWN"
                        parsed["step_magnitude_label_1_s"] = abs(change)
            subtype = "LAOS" if stype == "LAOS" else f"STEP_{parsed['step_direction'] or 'UNKNOWN'}"
            notes = []
            review_notes = []
            if not math.isfinite(median_dt):
                notes.append("No positive time increment")
                review_notes.append("No positive time increment")
            if np.any(dt <= 0):
                notes.append("Non-increasing time values")
                review_notes.append("Non-increasing time values")
            if math.isfinite(dt_cv) and dt_cv > 0.10:
                notes.append("Irregular sampling (dt CV > 0.10)")
            if stype == "STEP_UP_DOWN" and frame.shear_rate_1_s.notna().sum() < 12:
                notes.append("Shear-rate channel unavailable; direction relies on label")
            if stype == "STEP_UP_DOWN" and not parsed["step_direction"]:
                notes.append("Step direction unresolved")
                review_notes.append("Step direction unresolved")
            if stype == "STEP_UP_DOWN" and math.isfinite(parsed["step_final_rate_label_1_s"]) and math.isfinite(measured_final):
                rel = abs(measured_final - parsed["step_final_rate_label_1_s"]) / max(abs(parsed["step_final_rate_label_1_s"]), 0.1)
                if rel > 0.10:
                    notes.append("Measured final shear rate differs from label by >10%")
                    review_notes.append("Measured final shear rate differs from label by >10%")
            startup_outliers = 0
            if stype == "LAOS" and frame.strain_pct.notna().sum() >= 12:
                abs_strain = frame.strain_pct.abs()
                q995 = float(abs_strain.quantile(0.995))
                gamma_pct = parsed["strain_amplitude_label_pct"]
                limit = max(2.5 * gamma_pct, 1.25 * q995) if math.isfinite(gamma_pct) else 1.25 * q995
                startup_outliers = int((abs_strain > limit).sum())
                if startup_outliers:
                    notes.append(f"{startup_outliers} strain startup outlier row(s) flagged")
            record = {
                "source_file": path.name,
                "sheet": ws.title,
                "source_header_row": header_row,
                "source_start_column": col,
                "experiment_type": stype,
                "experiment_subtype": subtype,
                "condition_label": label,
                "n_points": len(frame),
                "duration_s": float(frame.time_s.max() - frame.time_s.min()),
                "median_dt_s": median_dt,
                "sampling_rate_hz": 1.0 / median_dt if math.isfinite(median_dt) and median_dt > 0 else math.nan,
                "dt_coefficient_of_variation": dt_cv,
                "has_stress": "YES",
                "has_strain": "YES" if frame.strain_pct.notna().sum() >= 12 else "NO",
                "has_shear_rate": "YES" if frame.shear_rate_1_s.notna().sum() >= 12 else "NO",
                "step_start_rate_measured_1_s": measured_start,
                "step_final_rate_measured_1_s": measured_final,
                "startup_outlier_rows": startup_outliers,
                "classification_confidence": "HIGH" if (stype == "LAOS" and math.isfinite(parsed["strain_amplitude_label_fraction"])) or (stype == "STEP_UP_DOWN" and parsed["step_direction"]) else "MEDIUM",
                "qc_status": "REVIEW" if review_notes else "OK",
                "qc_notes": "; ".join(notes),
                **parsed,
            }
            measurements.append(record)
            frame["source_file"] = path.name
            frame["sheet"] = ws.title
            frame["experiment_type"] = stype
            frame["experiment_subtype"] = subtype
            frame["condition_label"] = label
            frame["qc_flag"] = ""
            if startup_outliers:
                abs_strain = frame.strain_pct.abs()
                q995 = float(abs_strain.quantile(0.995))
                gamma_pct = parsed["strain_amplitude_label_pct"]
                limit = max(2.5 * gamma_pct, 1.25 * q995) if math.isfinite(gamma_pct) else 1.25 * q995
                frame.loc[abs_strain > limit, "qc_flag"] = "STARTUP_STRAIN_OUTLIER"
            signals.append(frame)
    return measurements, signals, issues, exclusions


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    file_rows, sheet_rows, measurement_rows, signal_frames, issue_rows, exclusion_rows = [], [], [], [], [], []
    for path in sorted(args.input_dir.glob("*.xlsx")):
        # Skipping Excel lock files and the separate DataZ cohort during core processing
        if path.name.startswith("~$") or path.name.casefold() == "dataz.xlsx":
            continue
        wb = load_workbook(path, read_only=False, data_only=False)
        sheet_meta = []
        for ws in wb.worksheets:
            meta = inspect_sheet(ws)
            meta["source_file"] = path.name
            sheet_meta.append(meta)
            sheet_rows.append(meta)
        measurements, signals, issues, exclusions = extract_core_blocks(path)
        measurement_rows.extend(measurements)
        signal_frames.extend(signals)
        issue_rows.extend(issues)
        exclusion_rows.extend(exclusions)
        file_rows.append(
            {
                "source_file": path.name,
                "size_bytes": path.stat().st_size,
                "size_mb": path.stat().st_size / (1024**2),
                "sheet_count": len(sheet_meta),
                "total_rows_across_sheets": sum(x["rows"] for x in sheet_meta),
                "total_nonempty_cells": sum(x["nonempty_cells"] for x in sheet_meta),
                "core_sheet_count": sum(x["part1_core_scope"] == "YES" for x in sheet_meta),
                "sheet_names": " | ".join(x["sheet"] for x in sheet_meta),
            }
        )

    files = pd.DataFrame(file_rows)
    sheets = pd.DataFrame(sheet_rows)
    measurements = pd.DataFrame(measurement_rows)
    signals = pd.concat(signal_frames, ignore_index=True) if signal_frames else pd.DataFrame()
    exclusions = pd.DataFrame(exclusion_rows)

    if len(measurements):
        measurements.insert(0, "measurement_id", [f"P1M{i:04d}" for i in range(1, len(measurements) + 1)])
        identified_frames = []
        for idx, frame in enumerate(signal_frames):
            identified = frame.copy()
            identified.insert(0, "measurement_id", measurements.iloc[idx].measurement_id)
            identified_frames.append(identified)
        signals = pd.concat(identified_frames, ignore_index=True)
        hashes = []
        for mid, group in signals.groupby("measurement_id", sort=False):
            hashes.append((mid, trace_hash(group)))
        hash_df = pd.DataFrame(hashes, columns=["measurement_id", "signal_sha256"])
        measurements = measurements.merge(hash_df, on="measurement_id", how="left")
        first_by_hash, duplicates = {}, []
        for _, row in measurements.iterrows():
            h = row.signal_sha256
            duplicates.append(first_by_hash.get(h, ""))
            first_by_hash.setdefault(h, row.measurement_id)
        measurements["duplicate_of_measurement_id"] = duplicates
        dup_mask = measurements.duplicate_of_measurement_id.ne("")
        measurements.loc[dup_mask, "qc_status"] = "REVIEW"
        measurements.loc[dup_mask, "qc_notes"] = measurements.loc[dup_mask, "qc_notes"].apply(
            lambda x: (str(x) + "; " if str(x) else "") + "Exact duplicate signal trace"
        )
        for _, row in measurements[measurements.qc_notes.fillna("").ne("")].iterrows():
            issue_rows.append(
                {
                    "measurement_id": row.measurement_id,
                    "source_file": row.source_file,
                    "sheet": row.sheet,
                    "condition_label": row.condition_label,
                    "severity": "REVIEW" if row.qc_status == "REVIEW" else "INFO",
                    "issue": row.qc_notes,
                }
            )
        signals = signals[[
            "measurement_id", "source_file", "sheet", "experiment_type", "experiment_subtype", "condition_label",
            "source_row", "time_s", "stress_pa", "strain_pct", "shear_rate_1_s", "qc_flag"
        ]]

    issues = pd.DataFrame(issue_rows)
    dictionary = pd.DataFrame(
        [
            ["measurement_index", "measurement_id", "Unique ID assigned to one detected measurement block", "text"],
            ["measurement_index", "experiment_type", "Core protocol class", "LAOS or STEP_UP_DOWN"],
            ["measurement_index", "experiment_subtype", "Detailed protocol class", "LAOS, STEP_UP, STEP_DOWN, STEP_UNKNOWN"],
            ["measurement_index", "n_points", "Aligned numeric time/stress rows in the block", "count"],
            ["measurement_index", "duration_s", "Last time minus first time", "s"],
            ["measurement_index", "sampling_rate_hz", "Inverse median positive time interval", "Hz"],
            ["measurement_index", "strain_amplitude_label_fraction", "Dimensionless LAOS gamma from source label", "fraction"],
            ["measurement_index", "strain_amplitude_label_pct", "LAOS gamma converted to percent", "%"],
            ["measurement_index", "step_direction", "Direction from source label or measured rate", "UP, DOWN, blank"],
            ["measurement_index", "duplicate_of_measurement_id", "First ID with identical numeric trace", "text or blank"],
            ["standardized_signals", "source_row", "Original worksheet row number", "row number"],
            ["standardized_signals", "time_s", "Elapsed time", "s"],
            ["standardized_signals", "stress_pa", "Measured shear stress", "Pa"],
            ["standardized_signals", "strain_pct", "Measured strain where available", "%"],
            ["standardized_signals", "shear_rate_1_s", "Measured shear rate where available", "1/s"],
            ["standardized_signals", "qc_flag", "Row-level quality flag; values are retained", "text or blank"],
        ],
        columns=["table", "column", "definition", "units_or_values"],
    )

    files.to_csv(args.output_dir / "file_inventory.csv", index=False)
    sheets.to_csv(args.output_dir / "sheet_inventory.csv", index=False)
    measurements.to_csv(args.output_dir / "measurement_index.csv", index=False)
    signals.to_csv(args.output_dir / "standardized_signals.csv.gz", index=False, compression="gzip")
    issues.to_csv(args.output_dir / "data_quality_log.csv", index=False)
    exclusions.to_csv(args.output_dir / "exclusion_log.csv", index=False)
    dictionary.to_csv(args.output_dir / "field_dictionary.csv", index=False)
    shutil.copy2(Path(__file__), args.output_dir / "data_processing_pipeline.py")

    counts = {
        "files": int(len(files)),
        "sheets": int(len(sheets)),
        "laos_measurements": int((measurements.experiment_subtype == "LAOS").sum()),
        "step_up_measurements": int((measurements.experiment_subtype == "STEP_UP").sum()),
        "step_down_measurements": int((measurements.experiment_subtype == "STEP_DOWN").sum()),
        "step_unknown_measurements": int((measurements.experiment_subtype == "STEP_UNKNOWN").sum()),
        "measurements_total": int(len(measurements)),
        "unique_signal_traces": int(measurements.duplicate_of_measurement_id.eq("").sum()),
        "standardized_signal_rows": int(len(signals)),
        "measurements_for_review": int((measurements.qc_status == "REVIEW").sum()),
        "excluded_or_catalog_only_records": int(len(exclusions)),
    }
    (args.output_dir / "processing_summary.json").write_text(json.dumps(counts, indent=2), encoding="utf-8")
    print(json.dumps(counts, indent=2))


if __name__ == "__main__":
    main()
